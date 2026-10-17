import json, zipfile, hashlib
from pathlib import Path
from lxml import etree
from docx import Document
p=Path(r'D:\career\VDT\GD2\upgrade_Ceph\MOP\MOP OSD.docx')
out=Path(r'D:\career\VDT\GD2\upgrade_Ceph\MOP\tmp\pa2-word')
ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
with zipfile.ZipFile(p) as z:
    root=etree.fromstring(z.read('word/document.xml'))
    texts=[]
    for el in root.findall('.//w:body/*', ns):
        if el.tag.endswith('}p'): texts.append(''.join(el.itertext()) if False else ''.join(el.xpath('.//w:t/text()',namespaces=ns)))
        elif el.tag.endswith('}tbl'):
            texts.append('[TABLE]')
            for row in el.findall('w:tr',ns): texts.append(' | '.join(''.join(c.xpath('.//w:t/text()',namespaces=ns)) for c in row.findall('w:tc',ns)))
    (out/'reference-text.txt').write_text('\n'.join(texts), encoding='utf-8')
    (out/'package-inventory.json').write_text(json.dumps([{'path':i.filename,'size':i.file_size,'sha256':hashlib.sha256(z.read(i.filename)).hexdigest()} for i in z.infolist()],indent=2),encoding='utf-8')
    (out/'document.xml').write_bytes(z.read('word/document.xml'))
    (out/'styles.xml').write_bytes(z.read('word/styles.xml'))
d=Document(p)
audit={'sha256':hashlib.sha256(p.read_bytes()).hexdigest(),'sections':[{'page_width':s.page_width.inches,'page_height':s.page_height.inches,'left_margin':s.left_margin.inches,'right_margin':s.right_margin.inches,'top_margin':s.top_margin.inches,'bottom_margin':s.bottom_margin.inches,'header_distance':s.header_distance.inches,'footer_distance':s.footer_distance.inches} for s in d.sections],'paragraphs':[{'idx':n,'style':p.style.name,'text':p.text,'runs':[{'font':r.font.name,'size':r.font.size.pt if r.font.size else None,'bold':r.bold} for r in p.runs]} for n,p in enumerate(d.paragraphs)], 'table_count':len(d.tables)}
(out/'reference-audit.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({k:v for k,v in audit.items() if k!='paragraphs'},ensure_ascii=False,indent=2))
print((out/'reference-text.txt').read_text(encoding='utf-8'))
